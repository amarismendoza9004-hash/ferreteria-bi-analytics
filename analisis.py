"""
Análisis: consultas SQL, detección de anomalías, gráficas y reporte en Excel.

Uso:
    python src/analisis.py
Requiere haber corrido antes src/etl.py.
"""
from __future__ import annotations

import re
import sqlite3
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.ticker as mtick
import numpy as np
import pandas as pd

RAIZ = Path(__file__).resolve().parents[1]
DB = RAIZ / "data" / "processed" / "ferreteria.db"
SQL = RAIZ / "sql" / "consultas_negocio.sql"
RES = RAIZ / "reports" / "resultados"
FIG = RAIZ / "reports" / "figures"
EXCEL = RAIZ / "reports" / "reporte_ferreteria.xlsx"

# Paleta categórica validada (orden fijo) + tintas de texto
SERIE = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#e34948"]
CRITICO = "#d03b3b"
TXT, TXT2, GRID, FONDO = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"

plt.rcParams.update({
    "figure.facecolor": FONDO, "axes.facecolor": FONDO, "savefig.facecolor": FONDO,
    "axes.edgecolor": GRID, "axes.labelcolor": TXT2, "xtick.color": TXT2, "ytick.color": TXT2,
    "axes.grid": True, "grid.color": GRID, "grid.linewidth": 0.8, "axes.axisbelow": True,
    "axes.spines.top": False, "axes.spines.right": False, "axes.spines.left": False,
    "font.size": 10, "axes.titlesize": 13, "axes.titleweight": "bold", "axes.titlecolor": TXT,
    "axes.titlelocation": "left", "legend.frameon": False, "lines.linewidth": 2,
})
PESOS = mtick.FuncFormatter(lambda x, _: f"${x/1e6:.1f} M" if abs(x) >= 1e6 else f"${x/1e3:,.0f} k")


# --------------------------------------------------------------------------- #
# SQL
# --------------------------------------------------------------------------- #
def leer_consultas(ruta: Path) -> dict[str, str]:
    """Separa el archivo .sql en consultas usando los comentarios '-- name:'."""
    texto = ruta.read_text(encoding="utf-8")
    bloques = re.split(r"^-- name:\s*(\w+)\s*$", texto, flags=re.M)
    return {bloques[i]: bloques[i + 1].strip().rstrip(";") for i in range(1, len(bloques), 2)}


def ejecutar_consultas(con: sqlite3.Connection) -> dict[str, pd.DataFrame]:
    RES.mkdir(parents=True, exist_ok=True)
    resultados = {}
    for nombre, query in leer_consultas(SQL).items():
        df = pd.read_sql_query(query, con)
        df.to_csv(RES / f"{nombre}.csv", index=False, encoding="utf-8-sig")
        resultados[nombre] = df
        print(f"  [SQL] {nombre}: {len(df)} filas")
    return resultados


# --------------------------------------------------------------------------- #
# Anomalías
# --------------------------------------------------------------------------- #
def _z_robusto(serie: pd.Series) -> tuple[pd.Series, pd.Series]:
    """Mediana móvil y z-score robusto (MAD) en una ventana centrada de 17 semanas."""
    r = serie.rolling(17, center=True, min_periods=6)
    mediana = r.median()
    mad = r.apply(lambda x: np.median(np.abs(x - np.median(x))), raw=True)
    return mediana, 0.6745 * (serie - mediana) / mad.replace(0, np.nan)


def anomalias_ventas_diarias(con: sqlite3.Connection, umbral_ventas: float = 5.0,
                             umbral_tickets: float = 4.0) -> pd.DataFrame:
    """
    Detecta días atípicos comparando cada día contra el MISMO día de la semana
    en las ~8 semanas previas y siguientes (mediana y MAD), para no confundir el
    sábado fuerte o el domingo flojo con anomalías.

    Se evalúan dos señales:
    - Importe vendido: detecta pedidos extraordinarios (picos).
    - Número de tickets: más estable que el importe, detecta caídas de
      operación (sistema POS caído, cortes de luz) aunque un solo ticket grande
      compense el importe.
    """
    d = pd.read_sql_query(
        "SELECT fecha_dia AS fecha, SUM(importe) AS ventas, COUNT(DISTINCT ticket_id) AS tickets "
        "FROM fact_ventas GROUP BY fecha_dia ORDER BY fecha_dia", con, parse_dates=["fecha"])
    d["dow"] = d["fecha"].dt.dayofweek
    partes = []
    for _, g in d.groupby("dow"):
        g = g.sort_values("fecha").copy()
        g["mediana"], g["z_robusto"] = _z_robusto(g["ventas"])
        g["tickets_esperados"], g["z_tickets"] = _z_robusto(g["tickets"])
        partes.append(g)
    d = pd.concat(partes).sort_values("fecha").reset_index(drop=True)
    d["anomalia"] = (d["z_robusto"].abs() > umbral_ventas) | (d["z_tickets"].abs() > umbral_tickets)
    d["tipo"] = np.select(
        [d["z_tickets"] < -umbral_tickets, d["z_robusto"] > umbral_ventas],
        ["Caída de operación", "Pico de venta"], "Variación atípica")
    return d


def resumen_anomalias(diarias: pd.DataFrame) -> pd.DataFrame:
    a = diarias[diarias["anomalia"]].copy()
    a["desviacion_vs_esperado_pct"] = (100 * (a["ventas"] / a["mediana"] - 1)).round(1)
    a = a[["fecha", "tipo", "ventas", "mediana", "desviacion_vs_esperado_pct", "tickets",
           "tickets_esperados", "z_robusto", "z_tickets"]].rename(columns={"mediana": "venta_esperada"})
    num = a.select_dtypes("number").columns
    a[num] = a[num].round(2)
    return a


# --------------------------------------------------------------------------- #
# Gráficas
# --------------------------------------------------------------------------- #
def guardar(fig, nombre: str) -> None:
    FIG.mkdir(parents=True, exist_ok=True)
    fig.savefig(FIG / f"{nombre}.png", dpi=150, bbox_inches="tight")
    plt.close(fig)


def graf_ventas_mensuales(df: pd.DataFrame) -> None:
    fig, ax = plt.subplots(figsize=(10, 4.2))
    x = pd.to_datetime(df["anio_mes"])
    ax.plot(x, df["ventas_netas"], color=SERIE[0], marker="o", markersize=4)
    ax.yaxis.set_major_formatter(PESOS)
    ax.set_ylim(0, df["ventas_netas"].max() * 1.15)
    ax.set_title("Ventas netas mensuales")
    for i, dy in ((df["ventas_netas"].idxmax(), 10), (df["ventas_netas"].idxmin(), -30)):
        ax.annotate(f"{df.loc[i, 'anio_mes']}\n${df.loc[i, 'ventas_netas']/1e6:.2f} M",
                    (x[i], df.loc[i, "ventas_netas"]), textcoords="offset points", xytext=(0, dy),
                    ha="center", fontsize=9, color=TXT2)
    guardar(fig, "01_ventas_mensuales")


def graf_categorias(df: pd.DataFrame) -> None:
    df = df.sort_values("ventas_netas")
    fig, ax = plt.subplots(figsize=(9, 4.8))
    ax.barh(df["categoria"], df["ventas_netas"], color=SERIE[0], height=0.6)
    ax.xaxis.set_major_formatter(PESOS)
    ax.grid(axis="y", visible=False)
    for y, (v, m) in enumerate(zip(df["ventas_netas"], df["margen_pct"])):
        ax.text(v, y, f"  ${v/1e6:.1f} M · margen {m:.0f}%", va="center", fontsize=9, color=TXT2)
    ax.set_xlim(0, df["ventas_netas"].max() * 1.35)
    ax.set_title("Ventas por categoría (jul 2024 – jun 2026)")
    guardar(fig, "02_ventas_categoria")


def graf_canales(con: sqlite3.Connection) -> None:
    d = pd.read_sql_query(
        "SELECT substr(fecha_dia,1,7) AS mes, canal, SUM(importe) AS ventas "
        "FROM fact_ventas GROUP BY mes, canal", con)
    p = d.pivot(index="mes", columns="canal", values="ventas")
    p = p.div(p.sum(axis=1), axis=0) * 100
    orden = ["Mostrador", "Mercado Libre", "WhatsApp", "Amazon"]
    fig, ax = plt.subplots(figsize=(10, 4.2))
    x = pd.to_datetime(p.index)
    # etiquetas al final de cada línea, separadas al menos 3.5 puntos para que no se encimen
    finales = sorted(((p[c].iloc[-1], c) for c in orden), reverse=True)
    pos, ultimo = {}, None
    for v, c in finales:
        y = v if ultimo is None else min(v, ultimo - 3.5)
        pos[c], ultimo = y, y
    for i, c in enumerate(orden):
        ax.plot(x, p[c], color=SERIE[i], label=c)
        ax.text(x[-1], pos[c], f"  {c} {p[c].iloc[-1]:.0f}%", va="center", fontsize=9, color=TXT2)
    ax.yaxis.set_major_formatter(mtick.PercentFormatter(decimals=0))
    ax.set_ylim(0, 80)
    ax.set_xlim(x[0], x[-1] + pd.Timedelta(days=110))
    ax.legend(loc="upper center", ncol=4, bbox_to_anchor=(0.45, -0.08))
    ax.set_title("Participación de cada canal en la venta mensual")
    guardar(fig, "03_participacion_canales")


def graf_estacionalidad(df: pd.DataFrame) -> None:
    p = df.pivot(index="categoria", columns="mes", values="indice_estacional")
    fig, ax = plt.subplots(figsize=(10, 4.6))
    # divergente: azul = debajo del promedio, naranja = arriba, gris = normal
    from matplotlib.colors import LinearSegmentedColormap, TwoSlopeNorm
    cmap = LinearSegmentedColormap.from_list("div", ["#2a78d6", "#efeeeb", "#eb6834"])
    im = ax.imshow(p.values, cmap=cmap, norm=TwoSlopeNorm(1.0, 0.4, 2.2), aspect="auto")
    ax.set_xticks(range(12), ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"])
    ax.set_yticks(range(len(p)), p.index)
    ax.grid(False)
    for (i, j), v in np.ndenumerate(p.values):
        ax.text(j, i, f"{v:.1f}", ha="center", va="center", fontsize=8,
                color="white" if abs(v - 1) > 0.7 else TXT)
    cb = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02)
    cb.outline.set_visible(False)
    cb.set_label("Índice (1.0 = mes promedio)", color=TXT2)
    ax.set_title("Estacionalidad por categoría")
    guardar(fig, "04_estacionalidad")


def graf_heatmap_horas(df: pd.DataFrame) -> None:
    dias = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    p = df.pivot(index="dia_semana", columns="hora", values="venta_promedio").reindex(dias)
    from matplotlib.colors import LinearSegmentedColormap
    cmap = LinearSegmentedColormap.from_list("seq", ["#eef4fc", "#2a78d6", "#0d3a73"])
    fig, ax = plt.subplots(figsize=(10, 3.6))
    im = ax.imshow(p.values, cmap=cmap, aspect="auto")
    ax.set_xticks(range(p.shape[1]), [f"{h}:00" for h in p.columns], fontsize=8)
    ax.set_yticks(range(7), dias)
    ax.grid(False)
    cb = fig.colorbar(im, ax=ax, fraction=0.025, pad=0.02, format=PESOS)
    cb.outline.set_visible(False)
    ax.set_title("Venta promedio en mostrador por día y hora")
    guardar(fig, "05_mapa_calor_horas")


def graf_anomalias(diarias: pd.DataFrame) -> None:
    fig, ax = plt.subplots(figsize=(11, 4.2))
    ax.plot(diarias["fecha"], diarias["ventas"], color=SERIE[0], linewidth=0.9, label="Venta diaria")
    a = diarias[diarias["anomalia"]]
    ax.scatter(a["fecha"], a["ventas"], s=60, color=CRITICO, edgecolor=FONDO, linewidth=2, zorder=3,
               label="Anomalía detectada")
    abajo = 0
    for r in a.itertuples():
        if r.ventas >= r.mediana:
            dy = 10
        else:
            dy, abajo = (-16 if abajo % 2 == 0 else -28), abajo + 1
        ax.annotate(f"{r.fecha:%d/%m/%y}", (r.fecha, r.ventas), textcoords="offset points",
                    xytext=(0, dy), ha="center", fontsize=8, color=TXT2)
    ax.yaxis.set_major_formatter(PESOS)
    ax.set_ylim(-diarias["ventas"].max() * 0.06, diarias["ventas"].max() * 1.12)
    ax.legend(loc="upper left")
    ax.set_title("Detección de días atípicos (z-score robusto por día de la semana)")
    guardar(fig, "06_anomalias_ventas")


def graf_inventario(df: pd.DataFrame) -> None:
    d = df.head(12).iloc[::-1]
    fig, ax = plt.subplots(figsize=(9, 4.6))
    etiquetas = [f"{s} · {t[:28]}" for s, t in zip(d["sku"], d["descripcion"])]
    ax.barh(etiquetas, d["valor_diferencia"], color=CRITICO, height=0.6)
    ax.xaxis.set_major_formatter(mtick.FuncFormatter(lambda x, _: f"-${abs(x):,.0f}" if x < 0 else f"${x:,.0f}"))
    ax.grid(axis="y", visible=False)
    for y, (v, u) in enumerate(zip(d["valor_diferencia"], d["diferencia"])):
        ax.text(v, y, f"{u:+d} pzas  ", va="center", ha="right", fontsize=8, color=TXT2)
    ax.set_xlim(d["valor_diferencia"].min() * 1.18, 0)
    ax.set_title("Faltantes de inventario con mayor impacto (conteo vs sistema)")
    guardar(fig, "07_faltantes_inventario")


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def reporte_excel(con: sqlite3.Connection, res: dict[str, pd.DataFrame], anomalias: pd.DataFrame) -> None:
    from openpyxl import load_workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    ventas = pd.read_sql_query(
        "SELECT v.fecha_dia, v.canal, p.categoria, v.importe, v.utilidad "
        "FROM fact_ventas v JOIN dim_producto p ON p.sku = v.sku", con, parse_dates=["fecha_dia"])
    ventas["anio_mes"] = ventas["fecha_dia"].dt.strftime("%Y-%m")
    tabla_din = pd.pivot_table(ventas, index="categoria", columns="canal", values="importe",
                               aggfunc="sum", margins=True, margins_name="Total").round(2)
    mensual_cat = pd.pivot_table(ventas, index="anio_mes", columns="categoria", values="importe",
                                 aggfunc="sum").round(2)

    hojas = {
        "Resumen KPIs": res["kpis_generales"],
        "Categoría x Canal": tabla_din.reset_index(),
        "Mensual x Categoría": mensual_cat.reset_index(),
        "Canales": res["crecimiento_semestral_canal"],
        "Top productos": res["top_productos"],
        "ABC": res["clasificacion_abc"],
        "Reabastecimiento": res["reabastecimiento"],
        "Dif. inventario": res["diferencias_inventario"],
        "Erosión margen": res["erosion_margen_proveedor"],
        "Precio < costo": res["productos_precio_bajo_costo"],
        "Anomalías": anomalias.assign(fecha=anomalias["fecha"].dt.date),
    }
    with pd.ExcelWriter(EXCEL, engine="openpyxl") as xw:
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre, index=False, startrow=1)

    wb = load_workbook(EXCEL)
    encabezado = PatternFill("solid", fgColor="1F4E78")
    for ws in wb.worksheets:
        ws["A1"] = ws.title
        ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
        for c in ws[2]:
            c.fill, c.font = encabezado, Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, col in enumerate(ws.iter_cols(min_row=2), start=1):
            ancho = max(len(str(c.value or "")) for c in col[:60])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, ancho + 2), 45)
            if any(k in str(col[0].value) for k in ("ventas", "utilidad", "importe", "costo", "precio",
                                                     "valor", "ene_jun", "esperada", "perdida")) \
               or ws.title in ("Categoría x Canal", "Mensual x Categoría") and i > 1:
                for c in col[1:]:
                    c.number_format = '"$"#,##0.00'
        ws.freeze_panes = "B3"
        ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"

    # Formato condicional: margen bajo en rojo, escala de color en tabla dinámica
    rojo = PatternFill("solid", fgColor="F8D7DA")
    for hoja, col in (("Erosión margen", "I"), ("Precio < costo", "F")):
        ws = wb[hoja]
        ws.conditional_formatting.add(f"{col}3:{col}{ws.max_row}",
                                      CellIsRule(operator="lessThan", formula=["20"], fill=rojo))
    ws = wb["Mensual x Categoría"]
    ws.conditional_formatting.add(f"B3:{get_column_letter(ws.max_column)}{ws.max_row}",
                                  ColorScaleRule(start_type="min", start_color="FFFFFF",
                                                 end_type="max", end_color="2A78D6"))

    # Hoja interactiva: selector con validación de datos + fórmulas
    ws = wb.create_sheet("Consulta rápida", 0)
    ws["A1"] = "Consulta rápida de ventas por categoría"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    ws["A3"], ws["B3"] = "Elige una categoría:", tabla_din.index[0]
    ws["B3"].fill = PatternFill("solid", fgColor="FFF2CC")
    n = len(tabla_din)
    dv = DataValidation(type="list", formula1=f"='Categoría x Canal'!$A$3:$A${n + 1}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B3")
    canales = [c for c in tabla_din.columns if c != "Total"]
    for i, canal in enumerate(canales + ["Total"], start=5):
        col = get_column_letter(list(tabla_din.columns).index(canal) + 2)
        ws[f"A{i}"] = canal
        ws[f"B{i}"] = f"=INDEX('Categoría x Canal'!{col}:{col},MATCH($B$3,'Categoría x Canal'!A:A,0))"
        ws[f"B{i}"].number_format = '"$"#,##0.00'
        ws[f"C{i}"] = f"=IFERROR(B{i}/$B${5 + len(canales)},0)"
        ws[f"C{i}"].number_format = "0.0%"
    ws[f"A{5 + len(canales)}"].font = ws[f"B{5 + len(canales)}"].font = Font(bold=True)
    ws["C4"], ws["B4"] = "% del total", "Ventas"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 24, 20
    wb.save(EXCEL)
    print(f"  [Excel] {EXCEL.relative_to(RAIZ)} ({len(wb.sheetnames)} hojas)")


def main() -> None:
    with sqlite3.connect(DB) as con:
        print("1/4 Ejecutando consultas SQL...")
        res = ejecutar_consultas(con)

        print("2/4 Detectando anomalías...")
        diarias = anomalias_ventas_diarias(con)
        anom = resumen_anomalias(diarias)
        anom.to_csv(RES / "anomalias_ventas_diarias.csv", index=False, encoding="utf-8-sig")
        print(f"  {len(anom)} días atípicos detectados")

        print("3/4 Generando gráficas...")
        graf_ventas_mensuales(res["ventas_mensuales"])
        graf_categorias(res["ventas_por_categoria"])
        graf_canales(con)
        graf_estacionalidad(res["estacionalidad_categoria"])
        graf_heatmap_horas(res["ventas_por_dia_hora"])
        graf_anomalias(diarias)
        graf_inventario(res["diferencias_inventario"])

        print("4/4 Construyendo reporte de Excel...")
        reporte_excel(con, res, anom)

    print("\nListo. Resultados en reports/")


if __name__ == "__main__":
    main()
